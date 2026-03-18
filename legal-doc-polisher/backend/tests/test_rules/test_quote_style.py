"""Tests for quote style rule."""

from docx import Document

from app.engine.deterministic.rules.quote_style import QuoteStyleRule
from app.engine.profiler import build_style_profile


def test_detects_straight_quotes(sample_dirty_docx):
    """Should detect straight quotes in the document."""
    profile = build_style_profile(sample_dirty_docx)
    doc = Document(str(sample_dirty_docx))
    rule = QuoteStyleRule()

    findings = rule.detect(doc, profile, {"target": "smart"})
    assert len(findings) >= 1


def test_fixes_straight_quotes(sample_dirty_docx):
    """Should convert straight quotes to smart quotes."""
    profile = build_style_profile(sample_dirty_docx)
    doc = Document(str(sample_dirty_docx))
    rule = QuoteStyleRule()

    findings = rule.detect(doc, profile, {"target": "smart"})
    for f in findings:
        rule.fix(doc, f, profile)

    # Verify no straight quotes remain in fixed paragraphs
    for f in findings:
        para = doc.paragraphs[f.paragraph_index]
        run = para.runs[f.run_index]
        assert '"' not in run.text
