"""Cross-reference and exhibit label consistency rule.

Common CRE legal doc issues:
- Broken cross-references ("Section 4.2" when section was renumbered to 4.3)
- Exhibit label mismatches ("Exhibit A" referenced but labeled "Exhibit B")
- Dangling references to sections/exhibits that don't exist
- Inconsistent reference formatting ("Section 3" vs "Sec. 3" vs "§3")
"""

from __future__ import annotations

import re
import uuid
from collections import defaultdict

from docx import Document

from ....models.enums import ChangeSource, Severity
from ....models.schemas import Change, Example, Finding, StyleProfile
from ....utils.docx_helpers import get_full_text, get_heading_level
from .base import Rule

# Patterns for section references in legal docs
SECTION_REF_PATTERNS = [
    re.compile(r"(?:Section|SECTION)\s+(\d+(?:\.\d+)*)", re.IGNORECASE),
    re.compile(r"(?:Sec\.|SEC\.)\s+(\d+(?:\.\d+)*)", re.IGNORECASE),
    re.compile(r"§\s*(\d+(?:\.\d+)*)"),
    re.compile(r"(?:Article|ARTICLE)\s+([IVXLCDM]+|\d+)", re.IGNORECASE),
]

# Patterns for exhibit/schedule/appendix references
EXHIBIT_REF_PATTERNS = [
    re.compile(r"(?:Exhibit|EXHIBIT)\s+([A-Z](?:-\d+)?)", re.IGNORECASE),
    re.compile(r"(?:Schedule|SCHEDULE)\s+([A-Z0-9](?:-\d+)?)", re.IGNORECASE),
    re.compile(r"(?:Appendix|APPENDIX)\s+([A-Z0-9])", re.IGNORECASE),
    re.compile(r"(?:Attachment|ATTACHMENT)\s+([A-Z0-9])", re.IGNORECASE),
    re.compile(r"(?:Addendum|ADDENDUM)\s+([A-Z0-9])", re.IGNORECASE),
]

# Patterns for definition references
DEFINED_TERM_REF = re.compile(r'"([A-Z][A-Za-z\s]+)"')

# Section reference format variants (for consistency checking)
SECTION_FORMAT_VARIANTS = [
    (re.compile(r"\bSection\s+\d"), "Section"),
    (re.compile(r"\bSec\.\s+\d"), "Sec."),
    (re.compile(r"\bSECTION\s+\d"), "SECTION"),
    (re.compile(r"§\s*\d"), "§"),
]


class CrossReferencesRule(Rule):
    name = "cross_references"
    category = "cross_references"

    def detect(
        self,
        document: Document,
        profile: StyleProfile,
        config: dict,
        examples: list[Example] | None = None,
    ) -> list[Finding]:
        findings = []

        # Build inventory of actual sections and exhibits
        section_numbers = self._collect_section_numbers(document)
        exhibit_labels = self._collect_exhibit_labels(document)

        # Scan for references and check validity
        findings.extend(self._check_section_refs(document, section_numbers))
        findings.extend(self._check_exhibit_refs(document, exhibit_labels))
        findings.extend(self._check_reference_format_consistency(document))

        return findings

    def _collect_section_numbers(self, document: Document) -> set[str]:
        """Build a set of all section numbers defined in the document."""
        sections = set()
        for para in document.paragraphs:
            level = get_heading_level(para)
            if level is None:
                continue
            text = get_full_text(para).strip()
            # Extract section number from heading text
            for pattern in SECTION_REF_PATTERNS:
                m = pattern.match(text)
                if m:
                    sections.add(m.group(1))
                    break
            # Also try to extract just leading numbers
            num_match = re.match(r"^(\d+(?:\.\d+)*)\s", text)
            if num_match:
                sections.add(num_match.group(1))
        return sections

    def _collect_exhibit_labels(self, document: Document) -> set[str]:
        """Build a set of all exhibit/schedule/appendix labels in the document."""
        labels = set()
        for para in document.paragraphs:
            text = get_full_text(para).strip().upper()
            for pattern in EXHIBIT_REF_PATTERNS:
                # Look for exhibit headings (typically standalone or at start of page)
                m = re.match(pattern.pattern, text, re.IGNORECASE)
                if m:
                    labels.add(m.group(1).upper())
        return labels

    def _check_section_refs(
        self, document: Document, known_sections: set[str]
    ) -> list[Finding]:
        """Check that all section references point to existing sections."""
        findings = []
        if not known_sections:
            return findings

        for idx, para in enumerate(document.paragraphs):
            if get_heading_level(para) is not None:
                continue
            text = get_full_text(para)
            for pattern in SECTION_REF_PATTERNS:
                for m in pattern.finditer(text):
                    ref = m.group(1)
                    # Check if the referenced section exists
                    if ref not in known_sections:
                        # Check partial matches (e.g., "4.2" might exist as part of "4.2.1")
                        partial = any(s.startswith(ref) or ref.startswith(s) for s in known_sections)
                        if not partial:
                            text_preview = text[:60]
                            findings.append(
                                Finding(
                                    paragraph_index=idx,
                                    category=self.category,
                                    location=f"Paragraph {idx + 1}: \"{text_preview}...\"",
                                    description=(
                                        f"Possible dangling cross-reference: "
                                        f"'{m.group(0)}' — section not found in document"
                                    ),
                                    severity=Severity.MAJOR,
                                    fix_data={"type": "dangling_section_ref", "ref": ref},
                                )
                            )
        return findings

    def _check_exhibit_refs(
        self, document: Document, known_exhibits: set[str]
    ) -> list[Finding]:
        """Check that all exhibit references match defined exhibits."""
        findings = []
        if not known_exhibits:
            return findings

        for idx, para in enumerate(document.paragraphs):
            text = get_full_text(para)
            for pattern in EXHIBIT_REF_PATTERNS:
                for m in pattern.finditer(text):
                    ref = m.group(1).upper()
                    if ref not in known_exhibits:
                        text_preview = text[:60]
                        findings.append(
                            Finding(
                                paragraph_index=idx,
                                category=self.category,
                                location=f"Paragraph {idx + 1}: \"{text_preview}...\"",
                                description=(
                                    f"Possible dangling exhibit reference: "
                                    f"'{m.group(0)}' — exhibit not found in document"
                                ),
                                severity=Severity.MAJOR,
                                fix_data={"type": "dangling_exhibit_ref", "ref": ref},
                            )
                        )
        return findings

    def _check_reference_format_consistency(self, document: Document) -> list[Finding]:
        """Check that section references use a consistent format throughout."""
        findings = []
        format_counts: dict[str, int] = defaultdict(int)
        format_locations: dict[str, list[int]] = defaultdict(list)

        for idx, para in enumerate(document.paragraphs):
            text = get_full_text(para)
            for pattern, fmt_name in SECTION_FORMAT_VARIANTS:
                count = len(pattern.findall(text))
                if count > 0:
                    format_counts[fmt_name] += count
                    format_locations[fmt_name].append(idx)

        if len(format_counts) > 1:
            # Find the dominant format
            dominant = max(format_counts, key=format_counts.get)
            for fmt_name, count in format_counts.items():
                if fmt_name != dominant and count > 0:
                    for idx in format_locations[fmt_name][:3]:  # Report first 3
                        text = get_full_text(document.paragraphs[idx])[:60]
                        findings.append(
                            Finding(
                                paragraph_index=idx,
                                category=self.category,
                                location=f"Paragraph {idx + 1}: \"{text}...\"",
                                description=(
                                    f"Inconsistent section reference format: "
                                    f"uses '{fmt_name}' but dominant format is '{dominant}'"
                                ),
                                severity=Severity.MODERATE,
                                fix_data={
                                    "type": "format_inconsistency",
                                    "current_format": fmt_name,
                                    "dominant_format": dominant,
                                },
                            )
                        )

        return findings

    def fix(self, document: Document, finding: Finding, profile: StyleProfile) -> Change:
        # Cross-reference fixes are mostly flagged for manual review
        # We can auto-fix format inconsistencies though
        fd = finding.fix_data

        if fd.get("type") == "format_inconsistency":
            para = document.paragraphs[finding.paragraph_index]
            current = fd.get("current_format", "")
            dominant = fd.get("dominant_format", "")

            if current and dominant:
                for run in para.runs:
                    # Replace format variant with dominant
                    if current == "Sec." and dominant == "Section":
                        run.text = re.sub(r"\bSec\.\s+", "Section ", run.text)
                    elif current == "SECTION" and dominant == "Section":
                        run.text = re.sub(r"\bSECTION\s+", "Section ", run.text)
                    elif current == "§" and dominant == "Section":
                        run.text = re.sub(r"§\s*", "Section ", run.text)

        return Change(
            change_id=str(uuid.uuid4()),
            category=self.category,
            location=finding.location,
            description=finding.description,
            severity=finding.severity,
            source=ChangeSource.DETERMINISTIC,
        )
