"""Word document parsing, section detection, and revision application."""

from __future__ import annotations

import copy
import os
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.text.paragraph import Paragraph


def parse_document(file_path: str) -> dict[str, Any]:
    """Parse a .docx file into structured sections with text content."""
    doc = Document(file_path)
    sections: list[dict[str, Any]] = []
    current_section: dict[str, Any] | None = None
    full_text_parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            if current_section:
                current_section["paragraphs"].append({"text": "", "style": para.style.name})
            full_text_parts.append("")
            continue

        full_text_parts.append(text)

        # Detect section headings (numbered sections, heading styles, ALL CAPS)
        is_heading = (
            para.style.name.startswith("Heading")
            or _is_numbered_heading(text)
            or (text.isupper() and len(text) < 80 and len(text.split()) < 10)
        )

        if is_heading:
            current_section = {
                "title": text,
                "heading_style": para.style.name,
                "paragraphs": [],
            }
            sections.append(current_section)
        else:
            if current_section is None:
                current_section = {
                    "title": "Preamble",
                    "heading_style": "Normal",
                    "paragraphs": [],
                }
                sections.append(current_section)
            current_section["paragraphs"].append({
                "text": text,
                "style": para.style.name,
            })

    # Extract tables
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            table_data.append([cell.text.strip() for cell in row.cells])
        tables.append(table_data)

    return {
        "sections": sections,
        "tables": tables,
        "full_text": "\n".join(full_text_parts),
        "paragraph_count": len(doc.paragraphs),
        "section_count": len(sections),
        "section_names": [s["title"] for s in sections],
    }


def apply_revisions(
    original_path: str,
    revisions: list[dict[str, Any]],
    output_path: str,
) -> str:
    """
    Apply approved revisions to a Word document.

    Each revision dict:
      {
        "original_text": "text to find",
        "revised_text": "replacement text",
        "section": "section name",
        "description": "what changed",
      }

    Returns the output file path.
    """
    doc = Document(original_path)

    for revision in revisions:
        original_text = revision["original_text"]
        revised_text = revision["revised_text"]

        for para in doc.paragraphs:
            if original_text in para.text:
                # Preserve formatting: replace text in runs
                _replace_in_paragraph(para, original_text, revised_text)
                break  # apply each revision once

    doc.save(output_path)
    return output_path


def _replace_in_paragraph(paragraph: Paragraph, old_text: str, new_text: str):
    """Replace text in a paragraph while trying to preserve run formatting."""
    # Build full text from runs
    full = "".join(run.text for run in paragraph.runs)
    if old_text not in full:
        return

    new_full = full.replace(old_text, new_text, 1)

    # Clear existing runs and set new text on first run, clear rest
    if paragraph.runs:
        # Keep first run's formatting, put all new text there
        paragraph.runs[0].text = new_full
        for run in paragraph.runs[1:]:
            run.text = ""


def generate_revision_summary(revisions: list[dict[str, Any]]) -> str:
    """Generate a human-readable summary of all revisions."""
    if not revisions:
        return "No revisions were made."

    lines = [f"## Revision Summary ({len(revisions)} changes)\n"]
    for i, rev in enumerate(revisions, 1):
        section = rev.get("section", "Unknown")
        desc = rev.get("description", "")
        lines.append(f"### {i}. {section}")
        lines.append(f"**Change:** {desc}")
        lines.append(f"- **Original:** {_truncate(rev.get('original_text', ''), 200)}")
        lines.append(f"- **Revised:** {_truncate(rev.get('revised_text', ''), 200)}")
        lines.append("")

    return "\n".join(lines)


def _is_numbered_heading(text: str) -> bool:
    """Check if text looks like a numbered section heading."""
    patterns = [
        r"^\d+\.\s+[A-Z]",       # "1. DEFINITIONS"
        r"^ARTICLE\s+\w+",       # "ARTICLE I"
        r"^SECTION\s+\d+",       # "SECTION 1"
        r"^[IVXLC]+\.\s+",      # "IV. ..."
        r"^\d+\.\d+\s+[A-Z]",   # "1.1 Definitions"
    ]
    return any(re.match(p, text) for p in patterns)


def _truncate(text: str, max_len: int) -> str:
    if len(text) <= max_len:
        return text
    return text[:max_len] + "..."
